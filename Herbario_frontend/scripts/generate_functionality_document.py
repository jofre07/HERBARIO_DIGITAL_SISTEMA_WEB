from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "documentacion"
OUT.mkdir(exist_ok=True)
OUTPUT = OUT / "Funcionalidad_del_Sistema_Herbario.docx"

GREEN = "245A37"
PALE = "EAF3EC"
PALE2 = "F6F8F6"
GRAY = "D9D9D9"
TEXT = RGBColor(36, 43, 38)

def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)

def set_cell_border(cell, color=GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), color)

def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + key))
        if node is None:
            node = OxmlElement("w:" + key)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True

def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    keep_with_next(p)
    return p

def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)

def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0].cells
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for i, header in enumerate(headers):
        hdr[i].text = header
        set_cell_fill(hdr[i], GREEN)
        for run in hdr[i].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
            run.font.size = Pt(9)
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx % 2:
                set_cell_fill(cells[i], PALE2)
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(9)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_border(cell)
            set_cell_margins(cell)
            if widths:
                cell.width = widths[i]
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.3)
section.right_margin = Cm(2.3)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = TEXT
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.12
for name, size in (("Title", 28), ("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)):
    st = styles[name]
    st.font.name = "Aptos Display" if name != "Normal" else "Aptos"
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
    st.paragraph_format.space_after = Pt(7)

# Portada
logo = ROOT / "public" / "images" / "logo-herbario.png"
if logo.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(logo), width=Inches(1.35))
    doc.inline_shapes[-1]._inline.docPr.set("descr", "Logotipo del Sistema Herbario")
    doc.inline_shapes[-1]._inline.docPr.set("title", "Sistema Herbario")
doc.add_paragraph().paragraph_format.space_after = Pt(28)
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Funcionalidad del Sistema Herbario")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("Descripción funcional del portal público y del panel administrativo")
r.font.size = Pt(15)
r.font.color.rgb = RGBColor(64, 83, 69)
doc.add_paragraph().paragraph_format.space_after = Pt(32)
meta = add_table(doc, ["Documento", "Detalle"], [
    ["Sistema", "Herbario frontend"],
    ["Tipo", "Documento funcional"],
    ["Versión analizada", "Aplicación Angular 19 del repositorio actual"],
    ["Fecha", "29 de septiembre de 2026"],
], [Cm(4.2), Cm(11.0)])
doc.add_paragraph().paragraph_format.space_after = Pt(35)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Este documento explica qué permite hacer el sistema, quién puede utilizar cada función y cómo se relacionan sus módulos principales.")
r.italic = True
r.font.size = Pt(10.5)
doc.add_page_break()

add_heading(doc, "1 Resumen ejecutivo", 1)
add_body(doc, "El Sistema Herbario es una aplicación web para registrar, organizar, administrar y divulgar información botánica. Su funcionamiento se divide en un portal público de consulta y un panel administrativo protegido. El portal facilita la exploración del catálogo, los especímenes, las colecciones, la galería y la distribución geográfica. El panel administrativo concentra el mantenimiento de datos, la gestión de usuarios y permisos, y la publicación de contenido.")
add_body(doc, "La solución aplica control de acceso mediante autenticación y permisos por módulo. Los datos se consumen desde servicios de una API, mientras que la interfaz ofrece búsqueda, filtros, paginación, formularios validados, confirmaciones y visualización cartográfica.")

add_heading(doc, "2 Objetivo y alcance", 1)
add_body(doc, "El objetivo del sistema es mantener un inventario digital consistente de información herbariológica y poner a disposición del público los registros autorizados para publicación. El alcance funcional observado incluye:")
add_bullets(doc, [
    "Consulta pública de especies, especímenes, imágenes, colecciones y ubicaciones.",
    "Administración de taxonomía, colectores, especímenes, imágenes y colecciones.",
    "Gestión de catálogos ambientales de climas y hábitats.",
    "Administración de usuarios, roles y permisos.",
    "Edición del contenido institucional de la página Acerca de.",
    "Indicadores estadísticos y mapas para apoyar el seguimiento del acervo.",
])

add_heading(doc, "3 Tipos de usuario", 1)
add_table(doc, ["Tipo de usuario", "Acceso", "Responsabilidades principales"], [
    ["Visitante", "Portal público", "Consultar información publicada, buscar registros, revisar imágenes, colecciones y mapas."],
    ["Usuario autenticado", "Panel administrativo", "Acceder al tablero y al perfil; las demás opciones dependen de sus permisos."],
    ["Gestor de datos", "Módulos autorizados", "Crear, consultar, actualizar o eliminar registros de los catálogos que tenga asignados."],
    ["Administrador", "Gestión de seguridad", "Administrar usuarios, roles, permisos y disponibilidad de funciones."],
], [Cm(3.2), Cm(3.4), Cm(8.6)])

add_heading(doc, "4 Estructura funcional", 1)
add_table(doc, ["Área", "Módulos"], [
    ["Portal público", "Inicio, catálogo, detalle de especie, especímenes, búsqueda avanzada, mapa, galería, colecciones y acerca de."],
    ["Administración", "Dashboard, perfil, usuarios, roles, climas, hábitats, taxonomía, colectores, especímenes, imágenes, colecciones y contenido institucional."],
    ["Servicios transversales", "Autenticación, token de sesión, permisos, validación de formularios, alertas, paginación, búsqueda y filtros."],
], [Cm(4.0), Cm(11.2)])

add_heading(doc, "5 Portal público", 1)
add_heading(doc, "5.1 Inicio", 2)
add_body(doc, "Presenta el acceso inicial a la plataforma y permite iniciar búsquedas que conducen al contenido botánico disponible. Funciona como punto de entrada hacia el catálogo y las principales secciones informativas.")

add_heading(doc, "5.2 Catálogo de especies", 2)
add_body(doc, "Muestra las especies publicadas y permite filtrar o buscar el catálogo. La navegación conduce a una ficha de detalle por especie, donde se reúne la información taxonómica y el material relacionado disponible.")

add_heading(doc, "5.3 Detalle de especie", 2)
add_body(doc, "La ficha consolida la identificación de la especie y sus datos asociados. También integra imágenes públicas y permite ampliar la galería o cargar más recursos visuales cuando existen registros adicionales.")

add_heading(doc, "5.4 Consulta de especímenes", 2)
add_body(doc, "Permite explorar los ejemplares publicados mediante búsqueda, filtros y paginación. El usuario puede abrir el detalle de un espécimen y consultar las imágenes relacionadas sin ingresar al panel administrativo.")

add_heading(doc, "5.5 Búsqueda avanzada", 2)
add_body(doc, "Ofrece filtros combinables para localizar registros con mayor precisión. Los filtros taxonómicos son dependientes: al elegir una familia se limitan los géneros disponibles, y al elegir un género se restringen las especies correspondientes.")

add_heading(doc, "5.6 Mapa de distribución", 2)
add_body(doc, "Representa ubicaciones georreferenciadas sobre un mapa Leaflet con cartografía de OpenStreetMap. Permite filtrar por familia, género, especie, nombre común y colector. Los resultados incluyen marcadores, fichas emergentes, acceso al detalle y estadísticas visibles de ubicaciones, especímenes, especies y colectores. El límite de la provincia de Manabí se resalta como referencia geográfica.")

add_heading(doc, "5.7 Galería y colecciones", 2)
add_body(doc, "La galería permite buscar y filtrar imágenes públicas, abrirlas en un visor ampliado y recorrer el material visual. La sección de colecciones presenta agrupaciones publicadas y admite búsqueda o filtrado para facilitar su exploración.")

add_heading(doc, "5.8 Información institucional", 2)
add_body(doc, "La página Acerca de presenta contenido institucional administrable. El portal incorpora encabezado, pie de página, navegación adaptable y un enlace para saltar al contenido principal, lo que favorece la navegación mediante teclado.")

add_heading(doc, "6 Panel administrativo", 1)
add_heading(doc, "6.1 Acceso y sesión", 2)
add_body(doc, "El ingreso se realiza mediante credenciales. Si no existe una sesión válida, el sistema redirige al inicio de sesión. Una vez autenticado, el usuario accede al dashboard y únicamente ve o abre las funciones permitidas por su perfil. El cierre de sesión requiere confirmación y elimina la sesión local.")

add_heading(doc, "6.2 Dashboard", 2)
add_body(doc, "Resume el estado del herbario mediante indicadores y gráficos. Presenta distribuciones por familia, año y estado de conservación, además de un mapa con los especímenes que poseen coordenadas válidas. Si una ubicación específica no está disponible, puede utilizar las coordenadas generales del espécimen.")

add_heading(doc, "6.3 Gestión de usuarios y roles", 2)
add_table(doc, ["Módulo", "Funciones"], [
    ["Usuarios", "Listar, buscar, paginar, crear, editar y eliminar usuarios; asignar roles y mantener el estado de la cuenta."],
    ["Roles", "Listar, buscar, crear, editar y eliminar roles; consultar y seleccionar permisos asociados."],
    ["Perfil", "Consultar la información del usuario autenticado desde el entorno administrativo."],
], [Cm(3.3), Cm(11.9)])

add_heading(doc, "6.4 Gestión taxonómica", 2)
add_body(doc, "La taxonomía se organiza en catálogos relacionados. Cada módulo dispone de consulta, búsqueda, paginación y operaciones de creación, edición y eliminación cuando el usuario cuenta con los permisos correspondientes.")
add_table(doc, ["Catálogo", "Propósito"], [
    ["Familias", "Mantener el nivel taxonómico superior utilizado para clasificar géneros y especies."],
    ["Géneros", "Registrar géneros y su relación con las familias."],
    ["Especies", "Administrar la identidad taxonómica que alimenta el catálogo público."],
    ["Estados", "Gestionar estados operativos o de clasificación usados por los registros."],
    ["Nombres comunes", "Relacionar denominaciones comunes con la información botánica."],
    ["Traducciones", "Mantener traducciones asociadas a los datos taxonómicos."],
    ["Estados de conservación", "Registrar categorías de conservación, incluidas las usadas en los indicadores."],
], [Cm(4.2), Cm(11.0)])

add_heading(doc, "6.5 Colectores", 2)
add_body(doc, "Administra la información de las personas responsables de la colecta. Los colectores se utilizan como referencia en especímenes, filtros públicos y puntos del mapa.")

add_heading(doc, "6.6 Especímenes y ubicaciones", 2)
add_body(doc, "Es el núcleo operativo del sistema. Permite registrar y mantener ejemplares, vincularlos con especies y colectores, controlar su publicación y asociar una o varias ubicaciones. El formulario incorpora pasos y validaciones para organizar la captura de información.")
add_bullets(doc, [
    "Código de catálogo e información descriptiva del ejemplar.",
    "Relaciones con especie, colector y otros catálogos disponibles.",
    "Una o varias ubicaciones, con identificación de la ubicación principal.",
    "Coordenadas geográficas y actualización visual del mapa desde los valores ingresados.",
    "Asociación de clima, hábitat y tipo de marcador para la representación cartográfica.",
    "Control de publicación para separar registros internos de los visibles al público.",
])

add_heading(doc, "6.7 Imágenes", 2)
add_body(doc, "Permite registrar imágenes del herbario, asociarlas con especímenes, previsualizarlas y mantener su información. El módulo admite selección de archivos, edición, eliminación y control de publicación. Las imágenes publicadas alimentan la galería y las fichas públicas.")

add_heading(doc, "6.8 Colecciones", 2)
add_body(doc, "Permite crear agrupaciones de especímenes, editar su contenido y definir su publicación. La interfaz carga especímenes disponibles para establecer las relaciones de cada colección. Las colecciones publicadas se muestran en el portal público.")

add_heading(doc, "6.9 Climas y hábitats", 2)
add_body(doc, "Estos catálogos ambientales se administran mediante un componente común configurado según el tipo de dato. Sus valores se relacionan con las ubicaciones de los especímenes y ayudan a describir el contexto ambiental de la colecta.")

add_heading(doc, "6.10 Contenido de la página Acerca de", 2)
add_body(doc, "Permite editar el contenido institucional que se publica en el portal. El administrador puede actualizar el texto y gestionar una imagen asociada, incluida su sustitución o eliminación.")

add_heading(doc, "7 Flujos principales", 1)
add_heading(doc, "7.1 Consulta pública", 2)
add_numbered(doc, [
    "El visitante ingresa al portal y selecciona catálogo, especímenes, búsqueda avanzada, mapa, galería o colecciones.",
    "El sistema solicita a la API únicamente la información publicada correspondiente a la vista.",
    "El visitante aplica búsquedas, filtros o paginación para reducir los resultados.",
    "El visitante abre una ficha, imagen o marcador para revisar información detallada.",
])

add_heading(doc, "7.2 Registro administrativo de un espécimen", 2)
add_numbered(doc, [
    "El usuario inicia sesión y accede al módulo si posee el permiso de consulta de especímenes.",
    "Abre el formulario de creación y completa los datos requeridos del ejemplar.",
    "Selecciona relaciones taxonómicas, colector y catálogos auxiliares.",
    "Registra una o varias ubicaciones, coordenadas y datos ambientales.",
    "El sistema valida el formulario y envía la información a la API.",
    "Tras una respuesta correcta, actualiza el listado y muestra una confirmación.",
])

add_heading(doc, "7.3 Publicación de información", 2)
add_numbered(doc, [
    "El gestor crea o edita el registro administrativo.",
    "Define si el registro debe estar publicado, cuando el módulo ofrece esta propiedad.",
    "Los servicios públicos consultan solo el contenido habilitado para publicación.",
    "El portal refleja la información en catálogo, especímenes, galería o colecciones según corresponda.",
])

add_heading(doc, "8 Reglas funcionales y validaciones", 1)
add_bullets(doc, [
    "Las rutas administrativas requieren autenticación.",
    "Cada módulo protegido valida un permiso específico antes de permitir el acceso.",
    "Las acciones visibles en tablas y formularios también pueden condicionarse por permiso.",
    "Los formularios aplican campos obligatorios y validadores de formato, incluyendo controles para evitar contenido HTML no permitido.",
    "Las operaciones sensibles, como eliminación o cierre de sesión, utilizan confirmación previa.",
    "Las coordenadas deben ser numéricas y válidas para que un registro se represente en el mapa.",
    "Los filtros taxonómicos mantienen coherencia entre familia, género y especie.",
    "La información no publicada queda disponible para la administración, pero se excluye de las consultas públicas.",
])

add_heading(doc, "9 Seguridad y control de acceso", 1)
add_body(doc, "La aplicación conserva el token y los datos del usuario mediante un servicio de sesión. Un interceptor adjunta la autenticación a las solicitudes correspondientes. El guard de autenticación protege el área administrativa y el guard de permisos comprueba la autorización requerida por cada ruta. Si el usuario no cumple el permiso, el sistema lo devuelve al dashboard.")
add_body(doc, "El control del frontend mejora la experiencia y evita accesos accidentales, pero la API debe validar nuevamente autenticación, permisos y reglas de publicación en cada operación. La autorización definitiva no debe depender únicamente de la interfaz.", "El control del frontend")

add_heading(doc, "10 Integraciones y tecnología observable", 1)
add_table(doc, ["Elemento", "Uso funcional"], [
    ["Angular 19", "Estructura de componentes independientes, rutas, formularios y consumo de servicios."],
    ["API HTTP", "Fuente y destino de los datos del herbario mediante servicios especializados."],
    ["RxJS", "Coordinación de solicitudes y manejo reactivo de respuestas."],
    ["Leaflet y OpenStreetMap", "Visualización de ubicaciones y distribución de especímenes."],
    ["Bootstrap y Bootstrap Icons", "Base visual y recursos de interfaz."],
    ["SweetAlert2", "Mensajes de éxito, error y confirmación."],
], [Cm(4.0), Cm(11.2)])

add_heading(doc, "11 Manejo de estados y errores", 1)
add_body(doc, "Las pantallas contemplan estados de carga, resultados vacíos y errores de consulta. Los listados utilizan componentes reutilizables para búsqueda, paginación, selección de tamaño de página, acciones y marcadores de carga. Después de crear, editar o eliminar, la interfaz actualiza los datos y comunica el resultado mediante alertas.")
add_body(doc, "En el dashboard, las solicitudes de estadísticas y especímenes se coordinan para construir una vista conjunta. En los mapas se omiten coordenadas no válidas y se manejan de forma independiente los errores del límite geográfico de Manabí, evitando que un fallo accesorio inutilice la vista principal.")

add_heading(doc, "12 Consideraciones operativas", 1)
add_bullets(doc, [
    "La calidad del catálogo depende de mantener relaciones taxonómicas coherentes antes de registrar especímenes.",
    "Las coordenadas y el estado de publicación determinan la aparición de datos en mapas y vistas públicas.",
    "Los permisos deben asignarse según las funciones reales de cada usuario y revisarse periódicamente.",
    "Las imágenes deben asociarse al espécimen correcto y publicarse únicamente cuando estén listas para difusión.",
    "La información institucional puede mantenerse sin modificar el código del portal mediante el módulo de contenido.",
])

add_heading(doc, "13 Matriz resumida de módulos", 1)
add_table(doc, ["Módulo", "Público", "Administración", "Operaciones principales"], [
    ["Especies", "Sí", "Sí", "Consulta pública; alta, edición y eliminación administrativa."],
    ["Especímenes", "Sí", "Sí", "Búsqueda y detalle; mantenimiento, ubicaciones y publicación."],
    ["Imágenes", "Sí", "Sí", "Galería y visor; carga, edición, asociación y publicación."],
    ["Colecciones", "Sí", "Sí", "Exploración; creación, edición, relación de especímenes y publicación."],
    ["Mapa", "Sí", "Dashboard", "Filtros, marcadores, estadísticas y referencia geográfica."],
    ["Taxonomía", "Indirecto", "Sí", "Catálogos maestros y relaciones taxonómicas."],
    ["Colectores", "Filtros", "Sí", "Consulta asociada y mantenimiento del catálogo."],
    ["Usuarios y roles", "No", "Sí", "Cuentas, roles y permisos."],
    ["Contenido institucional", "Lectura", "Sí", "Publicación y edición de la página Acerca de."],
], [Cm(3.0), Cm(2.1), Cm(2.8), Cm(7.3)])

add_heading(doc, "14 Conclusión", 1)
add_body(doc, "El Sistema Herbario cubre el ciclo funcional esencial de un repositorio botánico: estructura los catálogos de referencia, registra ejemplares y ubicaciones, vincula imágenes y colecciones, controla qué información se publica y ofrece herramientas públicas de consulta y visualización. La separación entre portal y administración, junto con el esquema de permisos, permite divulgar el acervo sin exponer las operaciones internas de mantenimiento.")

# Footer and page numbers
for sec in doc.sections:
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Sistema Herbario  |  Documento funcional  |  ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)
    for run in p.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 110, 103)

doc.core_properties.title = "Funcionalidad del Sistema Herbario"
doc.core_properties.subject = "Descripción funcional del portal público y del panel administrativo"
doc.core_properties.author = "Proyecto Herbario"
doc.core_properties.keywords = "herbario, funcionalidad, sistema, portal, administración"
doc.save(OUTPUT)
print(OUTPUT)
