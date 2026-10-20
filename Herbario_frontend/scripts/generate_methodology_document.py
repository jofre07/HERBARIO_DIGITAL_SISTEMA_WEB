from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "documentacion"
OUT.mkdir(exist_ok=True)
OUTPUT = OUT / "Metodologia_de_Desarrollo_del_Sistema_Herbario_Justificado.docx"

GREEN = "245A37"
PALE = "EEF5EF"
PALE2 = "F7F9F7"
GRAY = "D9D9D9"
TEXT = RGBColor(36, 43, 38)
TABLE_NUMBER = 0
FIGURE_NUMBER = 0


def fill(cell, color):
    pr = cell._tc.get_or_add_tcPr()
    node = pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        pr.append(node)
    node.set(qn("w:fill"), color)


def borders(cell):
    pr = cell._tc.get_or_add_tcPr()
    box = pr.first_child_found_in("w:tcBorders")
    if box is None:
        box = OxmlElement("w:tcBorders")
        pr.append(box)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = box.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            box.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), GRAY)


def margins(cell, value=120):
    pr = cell._tc.get_or_add_tcPr()
    box = pr.first_child_found_in("w:tcMar")
    if box is None:
        box = OxmlElement("w:tcMar")
        pr.append(box)
    for edge in ("top", "start", "bottom", "end"):
        node = box.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            box.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def body(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def bullets(doc, values):
    for value in values:
        p = doc.add_paragraph(value, style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def numbered(doc, values):
    for value in values:
        p = doc.add_paragraph(value, style="List Number")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_seq_caption(doc, label, title, number):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run(f"{label} ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f"SEQ {label} \\* ARABIC")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = str(number)
    run.append(text)
    field.append(run)
    p._p.append(field)
    p.add_run(f". {title}")
    return p


def add_index_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(17)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_field(doc, instruction, cached_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = cached_text
    run.append(text)
    field.append(run)
    p._p.append(field)
    return p


def table(doc, headers, rows, widths=None, caption=None):
    global TABLE_NUMBER
    if caption:
        TABLE_NUMBER += 1
        add_seq_caption(doc, "Tabla", caption, TABLE_NUMBER)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    header_pr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_pr.append(repeat)
    for i, text in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = text
        fill(cell, GREEN)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8.7)
    for row_index, values in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
            if row_index % 2:
                fill(cells[i], PALE2)
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(8.7)
    for row in t.rows:
        for i, cell in enumerate(row.cells):
            borders(cell)
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if widths:
                cell.width = widths[i]
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.1)
section.bottom_margin = Cm(1.9)
section.left_margin = Cm(2.3)
section.right_margin = Cm(2.3)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = TEXT
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.12
styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name, size in (("Title", 27), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
    style = styles[name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(7)

# Portada
logo = ROOT / "public" / "images" / "logo-herbario.png"
if logo.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(logo), width=Inches(1.35))
    shape._inline.docPr.set("descr", "Logotipo del Sistema Herbario")
    FIGURE_NUMBER += 1
    add_seq_caption(doc, "Figura", "Logotipo del Sistema Herbario", FIGURE_NUMBER)
doc.add_paragraph().paragraph_format.space_after = Pt(25)
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Metodología de Desarrollo del Sistema Herbario")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Marco de trabajo para planificar, construir, validar y mantener la solución")
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(64, 83, 69)
doc.add_paragraph().paragraph_format.space_after = Pt(28)
table(doc, ["Documento", "Detalle"], [
    ["Sistema", "Sistema Herbario"],
    ["Tipo", "Metodología de desarrollo de software"],
    ["Enfoque", "Ágil, incremental y basado en Scrum"],
    ["Versión", "2.0"],
    ["Fecha", "30 de septiembre de 2026"],
], [Cm(4.2), Cm(11.0)])
doc.add_paragraph().paragraph_format.space_after = Pt(28)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Plan metodológico para construir el sistema en un periodo total de dos semanas.")
r.italic = True
doc.add_page_break()

add_index_heading(doc, "Tabla de contenido")
add_field(doc, 'TOC \\o "1-3" \\h \\z \\u', "La tabla de contenido se actualizará automáticamente al abrir el documento en Word.")
doc.add_page_break()
add_index_heading(doc, "Índice de tablas")
add_field(doc, 'TOC \\h \\z \\c "Tabla"', "El índice de tablas se actualizará automáticamente al abrir el documento en Word.")
add_index_heading(doc, "Índice de figuras")
add_field(doc, 'TOC \\h \\z \\c "Figura"', "El índice de figuras se actualizará automáticamente al abrir el documento en Word.")
doc.add_page_break()

heading(doc, "1 Resumen ejecutivo")
body(doc, "La creación del Sistema Herbario se gestionará mediante un enfoque ágil e incremental basado en Scrum, adaptado a un plazo total de dos semanas. El trabajo se organizará en dos sprints de cinco días laborables, con requerimientos priorizados, entregables funcionales y validación frecuente por parte de los responsables del herbario. Debido al tiempo disponible, el alcance deberá concentrarse en las funciones esenciales definidas para la primera versión.")
body(doc, "La metodología orientará la construcción del sistema completo: frontend web, backend, API y base de datos. El alcance previsto considera el portal público, el panel administrativo, catálogo, especímenes, colecciones, imágenes, mapas, taxonomía, ubicaciones, recolectores, contenido institucional, usuarios, roles y permisos. También establece prácticas para integración, pruebas, seguridad, accesibilidad y despliegue de todos los componentes.")
body(doc, "La regla principal de entrega será que ninguna historia de usuario se considere terminada hasta que cumpla sus criterios de aceptación, pase las revisiones de código y pruebas correspondientes, mantenga la compatibilidad con los contratos de la API y cuente con evidencia de validación.")

heading(doc, "2 Objetivo y alcance")
heading(doc, "2.1 Objetivo general", 2)
body(doc, "Definir un proceso repetible y verificable para desarrollar y evolucionar el Sistema Herbario, asegurando calidad funcional, mantenibilidad técnica, seguridad, accesibilidad y alineación con las necesidades de los usuarios.")
heading(doc, "2.2 Objetivos específicos", 2)
bullets(doc, [
    "Convertir las necesidades del herbario en historias de usuario claras y priorizadas.",
    "Entregar incrementos utilizables al finalizar cada sprint.",
    "Mantener la trazabilidad entre requerimientos, cambios de código, pruebas y aprobaciones.",
    "Prevenir defectos mediante validaciones tempranas, revisión de código y automatización.",
    "Proteger la información y respetar los roles y permisos definidos por el sistema.",
    "Facilitar el mantenimiento mediante componentes reutilizables, servicios desacoplados, modelos de datos consistentes y documentación actualizada.",
])
heading(doc, "2.3 Alcance metodológico", 2)
body(doc, "La metodología se aplica al análisis, diseño, implementación, pruebas y despliegue del frontend, el backend, la API y la base de datos. Incluye la definición de contratos de servicios, modelos y relaciones, reglas de negocio, autenticación, autorización, validaciones, migraciones y mecanismos de respaldo. Las políticas institucionales y los cambios de alcance deberán aprobarse antes de su implementación.")

heading(doc, "3 Metodología seleccionada")
heading(doc, "3.1 Enfoque Scrum adaptado", 2)
body(doc, "Se utilizará Scrum como marco principal, adaptado al tamaño del equipo y al plazo disponible. El Product Backlog concentrará las necesidades del producto y el desarrollo se dividirá en dos sprints de una semana. El primer sprint establecerá la base técnica y las funciones principales; el segundo completará la integración, las funciones restantes, las pruebas y el despliegue. Cuando el equipo sea pequeño, una persona podrá asumir más de un rol, pero las responsabilidades de priorización, facilitación y construcción deberán mantenerse diferenciadas.")
heading(doc, "3.2 Principios de trabajo", 2)
bullets(doc, [
    "Valor primero: se implementan antes las funciones que aportan mayor utilidad o reducen mayor riesgo.",
    "Incrementos pequeños: cada sprint produce resultados demostrables y potencialmente desplegables.",
    "Retroalimentación frecuente: los usuarios revisan avances antes de que los cambios sean costosos.",
    "Calidad incorporada: pruebas, accesibilidad y seguridad forman parte del desarrollo y no una fase final aislada.",
    "Transparencia: el estado, los impedimentos, los riesgos y las decisiones quedan registrados.",
    "Mejora continua: cada retrospectiva produce acciones concretas para el siguiente sprint.",
])
heading(doc, "3.3 Ciclo de una iteración", 2)
numbered(doc, [
    "Refinar y priorizar las historias candidatas.",
    "Planificar el objetivo del sprint y seleccionar trabajo según la capacidad disponible.",
    "Diseñar e implementar en incrementos pequeños, integrando con frecuencia.",
    "Ejecutar pruebas y resolver defectos dentro del mismo sprint.",
    "Demostrar el incremento a las partes interesadas y registrar observaciones.",
    "Realizar la retrospectiva y acordar una mejora medible.",
])

heading(doc, "4 Roles y responsabilidades")
table(doc, ["Rol", "Responsabilidades"], [
    ["Product Owner", "Define la visión, prioriza el backlog, aclara necesidades y acepta o rechaza resultados según criterios acordados."],
    ["Scrum Master o coordinador", "Facilita eventos, elimina impedimentos, protege el foco del sprint y promueve la mejora del proceso."],
    ["Desarrollo frontend", "Construye las interfaces, formularios, navegación, estados visuales, accesibilidad y consumo de la API."],
    ["Desarrollo backend", "Implementa modelos, reglas de negocio, API, autenticación, permisos, validaciones y acceso a datos."],
    ["Responsable de datos", "Define relaciones, restricciones, migraciones, datos iniciales, respaldos e integridad de la base de datos."],
    ["Especialista del herbario", "Valida terminología, reglas taxonómicas, datos de colecta, publicación y utilidad de los flujos."],
    ["Responsable de calidad", "Define casos de prueba, conserva evidencias y verifica criterios funcionales y no funcionales."],
    ["Responsable técnico", "Cuida arquitectura, integración, revisión de código, seguridad y decisiones técnicas de alto impacto."],
    ["Usuario validador", "Ejecuta pruebas de aceptación sobre escenarios reales y comunica hallazgos antes de liberar."],
], [Cm(4.0), Cm(11.2)], "Roles y responsabilidades del proyecto")

heading(doc, "5 Artefactos y trazabilidad")
table(doc, ["Artefacto", "Contenido", "Responsable principal"], [
    ["Visión y alcance", "Problema, objetivos, usuarios, límites y resultados esperados.", "Product Owner"],
    ["Product Backlog", "Historias, defectos, mejoras técnicas y prioridades.", "Product Owner"],
    ["Sprint Backlog", "Trabajo seleccionado, tareas, responsables y estimaciones.", "Equipo"],
    ["Criterios de aceptación", "Condiciones verificables que determinan la aceptación.", "Product Owner y calidad"],
    ["Definición de terminado", "Controles mínimos comunes para toda entrega.", "Equipo"],
    ["Evidencias de prueba", "Casos, resultados, capturas o registros y defectos.", "Calidad"],
    ["Registro de decisiones", "Decisiones técnicas o funcionales, alternativas y consecuencias.", "Responsable técnico"],
    ["Incremento", "Software integrado y validado producido durante el sprint.", "Equipo"],
], [Cm(3.3), Cm(7.9), Cm(4.0)], "Artefactos para la gestión y trazabilidad")
heading(doc, "5.1 Formato de historia de usuario", 2)
body(doc, "Las necesidades se redactarán con el patrón: Como tipo de usuario, quiero capacidad, para obtener beneficio. Cada historia deberá incluir contexto, criterios de aceptación, dependencias, prioridad, estimación y referencias de diseño cuando correspondan.")
body(doc, "Ejemplo: Como gestor del herbario, quiero registrar un espécimen con su especie, recolector y ubicación, para conservar un registro digital completo y permitir su consulta autorizada.")
heading(doc, "5.2 Criterios de aceptación del ejemplo", 2)
bullets(doc, [
    "El formulario identifica con claridad los campos obligatorios.",
    "La especie y el recolector se seleccionan desde catálogos válidos.",
    "Las coordenadas aceptan únicamente valores numéricos dentro de rangos permitidos.",
    "El sistema comunica los errores sin borrar los datos ingresados.",
    "Tras guardar correctamente, el registro aparece en el listado administrativo.",
    "La visibilidad pública depende del estado de publicación y de los permisos aplicables.",
])

heading(doc, "6 Eventos de Scrum")
table(doc, ["Evento", "Frecuencia", "Resultado esperado"], [
    ["Refinamiento", "Una o dos veces por sprint", "Historias comprensibles, estimables y preparadas para planificación."],
    ["Planificación", "Inicio del sprint", "Objetivo, alcance seleccionado, capacidad y plan de trabajo."],
    ["Seguimiento diario", "15 minutos por día hábil", "Avance visible, coordinación e impedimentos identificados."],
    ["Revisión", "Fin del sprint", "Demostración del incremento y retroalimentación registrada."],
    ["Retrospectiva", "Después de la revisión", "Una o dos mejoras concretas con responsable y fecha."],
], [Cm(3.2), Cm(4.0), Cm(8.0)], "Eventos de Scrum durante las dos semanas")
body(doc, "Si el equipo trabaja de forma académica o parcial, la reunión diaria puede sustituirse por una actualización breve escrita, siempre que conserve visibilidad sobre lo realizado, lo siguiente y los impedimentos.")

heading(doc, "7 Proceso de desarrollo")
heading(doc, "7.1 Descubrimiento y análisis", 2)
bullets(doc, [
    "Identificar usuarios, objetivos, procesos actuales, fuentes de información y restricciones.",
    "Modelar flujos de consulta pública, administración y publicación.",
    "Documentar reglas taxonómicas, permisos, validaciones y dependencias entre catálogos.",
    "Separar requerimientos indispensables, importantes y opcionales.",
    "Definir criterios no funcionales de desempeño, seguridad, accesibilidad y compatibilidad.",
])
heading(doc, "7.2 Diseño", 2)
body(doc, "El diseño definirá una arquitectura por capas que separe interfaz, API, reglas de negocio y persistencia. Para el frontend se especificarán componentes, rutas, estados, comportamiento responsive y accesibilidad. Para el backend se establecerán módulos, endpoints, DTO, validaciones, permisos y manejo de errores. El diseño de datos incluirá entidades, relaciones, restricciones, migraciones y datos iniciales indispensables.")
heading(doc, "7.3 Implementación", 2)
body(doc, "El desarrollo se realizará por historias verticales pequeñas que incluyan, cuando corresponda, interfaz, endpoint, regla de negocio y persistencia. En el frontend la presentación permanecerá en componentes y el acceso a la API en servicios tipados. En el backend se separarán controladores, servicios, validaciones y repositorios o mecanismos de acceso a datos. Los contratos de la API se acordarán antes de programar ambos extremos y se integrarán continuamente.")
heading(doc, "7.4 Integración y revisión", 2)
body(doc, "Cada cambio deberá compilar, aplicar las migraciones previstas, pasar las pruebas disponibles y someterse a revisión antes de integrarse. La revisión comprobará claridad, ausencia de duplicación, reglas de negocio, consultas, manejo de errores, permisos, accesibilidad, adaptación responsive y coherencia entre frontend, API y base de datos.")
heading(doc, "7.5 Validación y entrega", 2)
body(doc, "El incremento se probará en un ambiente controlado y se demostrará con datos representativos. Los defectos que impidan cumplir criterios de aceptación se resolverán antes de cerrar la historia. La liberación se acompañará de una lista de cambios, instrucciones de despliegue, riesgos conocidos y procedimiento de reversión.")

heading(doc, "8 Arquitectura y prácticas técnicas")
heading(doc, "8.1 Frontend", 2)
table(doc, ["Práctica", "Aplicación prevista"], [
    ["Componentes enfocados", "Cada componente resuelve una responsabilidad de interfaz y delega el acceso a datos en servicios."],
    ["Tipado estricto", "Se utilizan interfaces y tipos explícitos; se evita any y se validan respuestas externas."],
    ["Reutilización", "Se aprovechan tablas, paginación, búsqueda, modales, selectores, alertas y utilidades compartidas."],
    ["Programación reactiva", "RxJS o Signals coordinan estado y solicitudes, evitando suscripciones innecesarias o sin limpieza."],
    ["Rutas y seguridad", "Las rutas administrativas conservan guardas de autenticación y permisos."],
    ["Formularios", "Se aplican validadores, labels, mensajes comprensibles y estados disabled o loading."],
    ["Rendimiento", "Se mantiene lazy loading, se evitan cálculos repetidos en templates y se optimizan listas e imágenes."],
    ["Estilos", "Se utilizarán SCSS, Bootstrap y un sistema visual común sin mezclar librerías innecesarias."],
], [Cm(3.6), Cm(11.6)], "Prácticas técnicas del frontend")
heading(doc, "8.2 Backend y API", 2)
table(doc, ["Práctica", "Aplicación prevista"], [
    ["Arquitectura por capas", "Separar controladores, servicios, reglas de negocio y acceso a datos."],
    ["API REST", "Definir rutas, métodos HTTP, códigos de estado y formatos de respuesta consistentes."],
    ["Contratos", "Documentar endpoints, DTO, validaciones y errores antes de integrar con el frontend."],
    ["Autenticación", "Emitir y validar sesiones o tokens de forma segura; proteger todas las operaciones privadas."],
    ["Autorización", "Comprobar roles y permisos en el servidor para cada recurso y acción."],
    ["Validación", "Rechazar datos incompletos, inválidos o no autorizados antes de ejecutar reglas de negocio."],
    ["Errores y registros", "Devolver errores controlados y conservar registros técnicos sin exponer información sensible."],
    ["Documentación", "Mantener una referencia verificable de la API y ejemplos de solicitudes y respuestas."],
], [Cm(3.6), Cm(11.6)], "Prácticas técnicas del backend y la API")
heading(doc, "8.3 Base de datos", 2)
table(doc, ["Práctica", "Aplicación prevista"], [
    ["Modelo de datos", "Definir entidades, atributos, relaciones y cardinalidades según las reglas del herbario."],
    ["Integridad", "Aplicar claves, restricciones, unicidad y reglas referenciales para evitar inconsistencias."],
    ["Migraciones", "Versionar la estructura y ejecutar cambios de manera repetible en cada ambiente."],
    ["Datos iniciales", "Preparar únicamente catálogos y cuentas indispensables para validar la primera versión."],
    ["Consultas", "Optimizar filtros, búsquedas y relaciones utilizadas por los flujos prioritarios."],
    ["Respaldo", "Definir copia, restauración y verificación antes del despliegue final."],
], [Cm(3.6), Cm(11.6)], "Prácticas técnicas de la base de datos")

heading(doc, "9 Estrategia de control de versiones")
bullets(doc, [
    "Mantener una rama principal estable y protegida.",
    "Crear ramas breves por historia, corrección o tarea técnica.",
    "Usar mensajes de commit descriptivos que indiquen intención y alcance.",
    "Evitar mezclar refactorizaciones extensas con cambios funcionales no relacionados.",
    "Integrar mediante solicitudes de cambio revisadas por otra persona cuando el equipo lo permita.",
    "Relacionar cada solicitud con su historia, criterios, pruebas y evidencia visual cuando aplique.",
    "Etiquetar las versiones liberadas y conservar notas de versión.",
])
heading(doc, "9.1 Lista mínima para revisión de código", 2)
bullets(doc, [
    "La solución responde al requerimiento y no introduce cambios fuera de alcance.",
    "No se exponen credenciales, tokens ni datos sensibles.",
    "Se mantienen el tipado, la separación de responsabilidades y la reutilización.",
    "Los estados de carga, vacío, error y éxito están contemplados, y los permisos se validan tanto en cliente como en servidor.",
    "Las migraciones, consultas y reglas de negocio mantienen la integridad de los datos.",
    "La interfaz funciona con teclado, conserva contraste y se adapta a tamaños de pantalla relevantes.",
    "Las pruebas y la documentación se actualizan cuando el cambio lo requiere.",
])

heading(doc, "10 Estrategia de pruebas")
table(doc, ["Nivel", "Objetivo", "Ejemplos"], [
    ["Unitarias", "Validar unidades aisladas del frontend y backend.", "Validadores, helpers, componentes, servicios, reglas de negocio y permisos."],
    ["Integración", "Comprobar colaboración entre frontend, API y base de datos.", "Persistencia, catálogos, filtros dependientes, autenticación y errores."],
    ["Funcionales", "Verificar flujos completos según historias.", "Login, registro de espécimen, publicación, búsqueda y consulta pública."],
    ["Interfaz", "Revisar comportamiento visual y estados.", "Tablas, formularios, modales, mapas, carga, vacío y error."],
    ["Responsive", "Confirmar uso en escritorio, laptop, tableta y móvil.", "Navegación, filtros, tablas, formularios y acciones."],
    ["Accesibilidad", "Garantizar operación perceptible y mediante teclado.", "Labels, foco, orden, contraste, nombres accesibles y mensajes."],
    ["Seguridad", "Detectar fallos de sesión, permisos y validación.", "Acceso directo a rutas, token inválido, contenido no permitido y roles."],
    ["Base de datos", "Verificar integridad, migraciones y consultas prioritarias.", "Relaciones, restricciones, duplicados, respaldo y restauración."],
    ["Aceptación", "Confirmar utilidad para usuarios reales.", "Escenarios ejecutados por responsables del herbario."],
], [Cm(2.6), Cm(5.4), Cm(7.2)], "Estrategia de pruebas del sistema")
heading(doc, "10.1 Gestión de defectos", 2)
body(doc, "Cada defecto deberá registrar pasos de reproducción, resultado esperado, resultado observado, ambiente, evidencia, severidad y relación con la historia. Los defectos críticos bloquean la liberación. Los defectos menores pueden programarse posteriormente si el Product Owner acepta el riesgo y queda constancia de la decisión.")

heading(doc, "11 Seguridad, privacidad y datos")
bullets(doc, [
    "La API debe validar autenticación, permisos y reglas de negocio; la protección del frontend no sustituye la autorización del servidor.",
    "Los tokens y datos de sesión se manipulan únicamente mediante los servicios previstos y nunca se registran en texto plano.",
    "Toda entrada se valida tanto en cliente como en servidor, especialmente identificadores, coordenadas, archivos y contenido editable.",
    "Los permisos se aplican por mínimo privilegio y se revisan cuando cambian las funciones de un usuario.",
    "Las imágenes y datos se publican únicamente cuando su estado y autorización lo permiten.",
    "Los ambientes utilizan configuración separada y no incluyen secretos dentro del repositorio.",
    "Se definen copias de seguridad, restauración y retención para los datos administrados por el backend.",
    "Las consultas utilizan parámetros seguros y el backend limita la información expuesta por cada respuesta.",
])

heading(doc, "12 Accesibilidad y diseño responsive")
body(doc, "Toda nueva interfaz o modificación visual deberá funcionar en escritorio, laptop, tableta y móvil. Se utilizará HTML semántico, labels asociados, navegación por teclado, foco visible, contraste suficiente, mensajes de error claros y atributos ARIA únicamente cuando la semántica nativa no sea suficiente.")
table(doc, ["Aspecto", "Criterio de verificación"], [
    ["Navegación", "Todas las acciones esenciales son alcanzables y operables mediante teclado."],
    ["Formularios", "Cada control tiene label, instrucciones necesarias y error asociado."],
    ["Foco", "El indicador es visible y el orden coincide con la lectura de la pantalla."],
    ["Contraste", "Texto, controles e indicadores conservan legibilidad en todos sus estados."],
    ["Responsive", "No existe desplazamiento horizontal innecesario ni pérdida de acciones esenciales."],
    ["Estados", "Loading, disabled, vacío, éxito y error se comunican de manera perceptible."],
], [Cm(3.5), Cm(11.7)], "Criterios de accesibilidad y diseño responsive")

heading(doc, "13 Gestión de riesgos")
table(doc, ["Riesgo", "Probabilidad", "Impacto", "Respuesta"], [
    ["Cambios frecuentes de alcance", "Media", "Alto", "Priorizar backlog, congelar el sprint y evaluar cambios para iteraciones siguientes."],
    ["Contratos de API inestables", "Media", "Alto", "Versionar contratos, coordinar cambios y usar datos simulados para desacoplar trabajo."],
    ["Datos taxonómicos inconsistentes", "Media", "Alto", "Aplicar validaciones, catálogos controlados y revisión del especialista."],
    ["Permisos mal configurados", "Media", "Alto", "Matriz de roles, pruebas por perfil y mínimo privilegio."],
    ["Baja disponibilidad de validadores", "Media", "Medio", "Programar revisiones con anticipación y usar criterios explícitos y evidencias."],
    ["Defectos responsive o accesibles", "Media", "Medio", "Incluirlos en la definición de terminado y probar durante cada sprint."],
    ["Dependencia de una sola persona", "Media", "Alto", "Revisión cruzada, documentación y rotación de responsabilidades."],
    ["Pérdida o corrupción de datos", "Baja", "Crítico", "Respaldos verificados, restauración ensayada y controles de integridad."],
], [Cm(4.2), Cm(2.2), Cm(2.0), Cm(6.8)], "Matriz de riesgos del proyecto")

heading(doc, "14 Plan de trabajo de dos semanas")
body(doc, "La creación del sistema se ejecutará en diez días laborables. Para cumplir el plazo, el alcance de la primera versión se cerrará al inicio y cualquier solicitud nueva se registrará como mejora posterior, salvo que sustituya una función ya planificada.")
table(doc, ["Periodo", "Actividades principales", "Resultado"], [
    ["Día 1", "Acordar objetivo, alcance mínimo, roles, riesgos, historias y criterios de aceptación.", "Backlog priorizado y plan del Sprint 1."],
    ["Día 2", "Definir arquitectura completa, modelo de datos, navegación, contratos de API y diseño base.", "Arquitectura, modelo relacional, endpoints y prototipos esenciales."],
    ["Días 3 y 4", "Construir base de datos, migraciones, autenticación, permisos, API base, estructura frontend y catálogos fundamentales.", "Primer incremento de extremo a extremo."],
    ["Día 5", "Completar pruebas del Sprint 1, demostrar el avance y ajustar el backlog.", "Incremento revisado y plan del Sprint 2."],
    ["Días 6 y 7", "Completar endpoints, reglas de negocio e interfaces para especímenes, ubicaciones, imágenes, colecciones y consulta pública.", "Flujos prioritarios integrados con persistencia."],
    ["Día 8", "Completar validaciones cliente y servidor, permisos, responsive, accesibilidad, consultas y manejo de errores.", "Versión completa candidata."],
    ["Día 9", "Probar frontend, API y base de datos; ejecutar regresión, seguridad, migraciones y aceptación; corregir defectos prioritarios.", "Versión candidata validada de extremo a extremo."],
    ["Día 10", "Respaldar, desplegar frontend, backend y base de datos; verificar y entregar documentación técnica y funcional.", "Primera versión completa desplegada y aceptada."],
], [Cm(2.4), Cm(8.5), Cm(4.3)], "Cronograma de trabajo de diez días")
heading(doc, "14.1 Distribución por sprints", 2)
body(doc, "El Sprint 1 comprenderá los días 1 al 5 y dejará disponible una base integrada de frontend, backend y base de datos con los primeros flujos utilizables. El Sprint 2 abarcará los días 6 al 10 y se enfocará en completar los flujos prioritarios, estabilizar todos los componentes y preparar la entrega. Las pruebas se realizarán durante ambos sprints y no se concentrarán únicamente al final.")

heading(doc, "15 Definición de preparado y terminado")
heading(doc, "15.1 Definición de preparado", 2)
bullets(doc, [
    "La historia expresa usuario, necesidad y beneficio.",
    "Los criterios de aceptación son claros y comprobables.",
    "Las dependencias y datos necesarios están identificados.",
    "Existe diseño o referencia suficiente cuando la interfaz cambia.",
    "El equipo puede estimarla y completarla dentro de un sprint.",
])
heading(doc, "15.2 Definición de terminado", 2)
bullets(doc, [
    "El código está integrado, revisado y compila sin errores.",
    "Los criterios de aceptación han sido comprobados.",
    "Las pruebas aplicables se ejecutan con resultado satisfactorio.",
    "Se verifican autenticación, permisos, errores y estados de interfaz.",
    "Los endpoints, reglas de negocio, migraciones y restricciones de datos están probados.",
    "La vista es responsive, accesible mediante teclado y coherente con el sistema visual.",
    "No se alteran contratos de API sin aprobación y coordinación.",
    "La documentación y las notas de versión están actualizadas.",
    "El Product Owner o usuario designado acepta el resultado.",
])

heading(doc, "16 Indicadores de seguimiento")
table(doc, ["Indicador", "Propósito", "Frecuencia"], [
    ["Cumplimiento del objetivo del sprint", "Medir si el incremento logró el resultado acordado.", "Cada sprint"],
    ["Historias aceptadas", "Observar capacidad de entrega real, sin premiar volumen aislado.", "Cada sprint"],
    ["Defectos por severidad", "Detectar áreas inestables y priorizar prevención.", "Semanal"],
    ["Tiempo de resolución", "Controlar respuesta ante incidencias relevantes.", "Mensual"],
    ["Cobertura de pruebas críticas", "Asegurar protección de flujos esenciales.", "Cada versión"],
    ["Incidencias posteriores a liberación", "Evaluar efectividad de validación y regresión.", "Cada versión"],
    ["Satisfacción del usuario validador", "Confirmar utilidad y claridad de los flujos entregados.", "Cada revisión"],
], [Cm(4.5), Cm(7.4), Cm(3.3)], "Indicadores de seguimiento del desarrollo")
body(doc, "Los indicadores se utilizarán para aprender y ajustar el proceso. No deberán convertirse en metas aisladas que incentiven cerrar trabajo incompleto o reducir la calidad.")

heading(doc, "17 Despliegue y mantenimiento")
heading(doc, "17.1 Preparación del despliegue", 2)
numbered(doc, [
    "Confirmar la versión, los cambios incluidos y los criterios de aceptación.",
    "Ejecutar compilaciones de producción, migraciones controladas y pruebas de regresión priorizadas.",
    "Verificar configuración del frontend, backend, base de datos, URL de API y manejo de secretos.",
    "Preparar respaldo, plan de reversión y responsables de la ventana.",
    "Publicar y realizar pruebas rápidas de autenticación, consulta y operaciones críticas.",
    "Comunicar el resultado y registrar cualquier incidencia.",
])
heading(doc, "17.2 Mantenimiento", 2)
body(doc, "Las incidencias se clasificarán por severidad y urgencia. Las correcciones críticas seguirán un flujo abreviado, pero mantendrán revisión, pruebas y trazabilidad. Las mejoras evolutivas ingresarán al Product Backlog y competirán por prioridad con el resto del trabajo. Periódicamente se revisarán dependencias, accesibilidad, rendimiento, seguridad y deuda técnica.")

heading(doc, "18 Criterios de aceptación de la metodología")
bullets(doc, [
    "Existe un backlog único, priorizado y visible.",
    "Cada sprint tiene objetivo, capacidad y trabajo seleccionado.",
    "Las historias incluyen criterios verificables antes de comenzar.",
    "Cada entrega conserva evidencia de revisión y pruebas.",
    "Los usuarios del herbario participan en revisiones y aceptación.",
    "Los riesgos y decisiones relevantes se registran y actualizan.",
    "La definición de terminado se aplica de forma consistente.",
    "La retrospectiva produce acciones de mejora con seguimiento.",
])

heading(doc, "19 Conclusión")
body(doc, "La metodología propuesta organiza la creación completa del Sistema Herbario alrededor de entregas pequeñas, validación frecuente y controles de calidad integrados. Scrum aporta un ritmo predecible y visibilidad del trabajo, mientras que las prácticas técnicas definidas protegen la arquitectura del frontend y backend, los contratos de la API, la integridad de la base de datos, la seguridad, la accesibilidad y la experiencia de usuario.")
body(doc, "Su aplicación no depende de un equipo grande. Para cumplir el plazo de dos semanas será indispensable mantener un alcance mínimo estable, ejecutar dos sprints de una semana, validar diariamente los avances y aplicar una definición de terminado que impida considerar completa una función sin las comprobaciones necesarias.")

heading(doc, "Anexo A Plantilla breve de historia de usuario")
table(doc, ["Campo", "Contenido"], [
    ["Título", "Nombre breve y orientado al resultado."],
    ["Historia", "Como usuario, quiero capacidad, para beneficio."],
    ["Contexto", "Situación, regla o problema que origina la necesidad."],
    ["Criterios", "Condiciones observables para aceptar la historia."],
    ["Dependencias", "Servicios, datos, permisos, diseños u otras historias."],
    ["Pruebas", "Escenarios principales, alternos y de error."],
    ["Evidencia", "Enlace al cambio, resultados de prueba y aprobación."],
], [Cm(3.5), Cm(11.7)], "Plantilla de historia de usuario")

heading(doc, "Anexo B Lista de comprobación para liberación")
bullets(doc, [
    "Alcance y versión confirmados.",
    "Compilación de producción satisfactoria.",
    "Pruebas críticas y regresión aprobadas.",
    "Migraciones, integridad de datos, respaldo y restauración verificados.",
    "Permisos y autenticación verificados.",
    "Responsive y accesibilidad revisados.",
    "Configuración del ambiente confirmada.",
    "Respaldo y reversión preparados.",
    "Notas de versión y responsables comunicados.",
    "Pruebas posteriores al despliegue completadas.",
])

for sec in doc.sections:
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Sistema Herbario  |  Metodología de desarrollo  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)
    for run in p.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 110, 103)

doc.core_properties.title = "Metodología de Desarrollo del Sistema Herbario"
doc.core_properties.subject = "Proceso ágil para construir frontend, backend y base de datos del Sistema Herbario"
doc.core_properties.author = "Proyecto Herbario"
doc.core_properties.keywords = "herbario, metodología, Scrum, Angular, desarrollo de software"
settings = doc.settings._element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields")
    settings.append(update)
update.set(qn("w:val"), "true")
doc.save(OUTPUT)
print(OUTPUT)
